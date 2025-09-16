import pandas as pd
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

class ConfigManager:
    def __init__(self):
        self.config_file = 'RouteOptima_Config.xlsx'
        self._init_config_file()

    def _init_config_file(self):
        if not os.path.exists(self.config_file):
            self._create_config_file()
        else:
            # Check if all required sheets exist
            try:
                xl_file = pd.ExcelFile(self.config_file)
                required_sheets = ['Store_Locations', 'Vehicle_Config', 'Delivery_Locations', 'Company_Settings']
                missing_sheets = [sheet for sheet in required_sheets if sheet not in xl_file.sheet_names]
                if missing_sheets:
                    self._add_missing_sheets(missing_sheets)
            except:
                self._create_config_file()

    def _create_config_file(self):
        wb = Workbook()
        
        # Store Locations Sheet
        ws_stores = wb.active
        ws_stores.title = "Store_Locations"
        store_headers = ['Store_ID', 'Store_Name', 'Latitude', 'Longitude', 'Is_Active']
        for col, header in enumerate(store_headers, 1):
            cell = ws_stores.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        
        # Sample store data
        sample_stores = [
            [1, "Main Store", 23.0225, 72.5714, "Yes"],
            [2, "Branch Store", 23.0325, 72.5814, "Yes"]
        ]
        for row, data in enumerate(sample_stores, 2):
            for col, value in enumerate(data, 1):
                ws_stores.cell(row=row, column=col, value=value)
        
        # Vehicle Configuration Sheet
        ws_vehicles = wb.create_sheet("Vehicle_Config")
        vehicle_headers = ['Vehicle_ID', 'Store_ID', 'Vehicle_Type', 'Count', 'Capacity_Shipments', 'Max_Radius_KM', 'Max_Trip_Time_Minutes', 'Is_Active']
        for col, header in enumerate(vehicle_headers, 1):
            cell = ws_vehicles.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
        
        # Sample vehicle data
        sample_vehicles = [
            [1, 1, "3W", 50, 5, 15, 240, "Yes"],
            [2, 1, "4W-EV", 25, 8, 20, 300, "Yes"],
            [3, 1, "4W", "Unlimited", 25, "Unlimited", 480, "Yes"]
        ]
        for row, data in enumerate(sample_vehicles, 2):
            for col, value in enumerate(data, 1):
                ws_vehicles.cell(row=row, column=col, value=value)
        
        # Delivery Locations Sheet
        ws_deliveries = wb.create_sheet("Delivery_Locations")
        delivery_headers = ['Delivery_ID', 'Store_ID', 'Customer_Name', 'Address', 'Latitude', 'Longitude', 'Delivery_Timeslot', 'Is_Active']
        for col, header in enumerate(delivery_headers, 1):
            cell = ws_deliveries.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="F39C12", end_color="F39C12", fill_type="solid")
        
        # Sample delivery data
        sample_deliveries = [
            [1001, 1, "John Doe", "123 Main St", 23.0325, 72.5814, "09:00-12:00", "Yes"],
            [1002, 1, "Jane Smith", "456 Oak Ave", 23.0225, 72.5914, "09:00-12:00", "Yes"],
            [1003, 1, "Bob Johnson", "789 Pine Rd", 23.0425, 72.5714, "12:00-15:00", "Yes"],
            [1004, 1, "Alice Brown", "321 Elm St", 23.0125, 72.5614, "15:00-18:00", "Yes"]
        ]
        for row, data in enumerate(sample_deliveries, 2):
            for col, value in enumerate(data, 1):
                ws_deliveries.cell(row=row, column=col, value=value)
        
        # Company Settings Sheet
        ws_settings = wb.create_sheet("Company_Settings")
        settings_headers = ['Setting_Name', 'Setting_Value', 'Description']
        for col, header in enumerate(settings_headers, 1):
            cell = ws_settings.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid")
        
        # Sample settings
        sample_settings = [
            ["Company_Name", "Your Company Name", "Name of your company"],
            ["Default_Store_ID", "1", "Default store for route planning"],
            ["Working_Hours_Start", "09:00", "Business start time"],
            ["Working_Hours_End", "18:00", "Business end time"]
        ]
        for row, data in enumerate(sample_settings, 2):
            for col, value in enumerate(data, 1):
                ws_settings.cell(row=row, column=col, value=value)
        
        wb.save(self.config_file)
    
    def _add_missing_sheets(self, missing_sheets):
        from openpyxl import load_workbook
        wb = load_workbook(self.config_file)
        
        if 'Delivery_Locations' in missing_sheets:
            ws = wb.create_sheet('Delivery_Locations')
            headers = ['Delivery_ID', 'Store_ID', 'Customer_Name', 'Address', 'Latitude', 'Longitude', 'Delivery_Timeslot', 'Is_Active']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="F39C12", end_color="F39C12", fill_type="solid")
        
        if 'Vehicle_Config' in missing_sheets:
            # Update existing Vehicle_Config to include Store_ID
            if 'Vehicle_Config' in wb.sheetnames:
                ws = wb['Vehicle_Config']
                # Check if Store_ID column exists
                if ws.cell(1, 2).value != 'Store_ID':
                    # Insert Store_ID column
                    ws.insert_cols(2)
                    ws.cell(1, 2, 'Store_ID')
                    ws.cell(1, 2).font = Font(bold=True, color="FFFFFF")
                    ws.cell(1, 2).fill = PatternFill(start_color="27AE60", end_color="27AE60", fill_type="solid")
                    # Add Store_ID = 1 for existing vehicles
                    for row in range(2, ws.max_row + 1):
                        ws.cell(row, 2, 1)
        
        wb.save(self.config_file)

    def get_stores(self):
        df = pd.read_excel(self.config_file, sheet_name='Store_Locations')
        return df[df['Is_Active'] == 'Yes'].to_dict('records')

    def get_vehicles(self):
        df = pd.read_excel(self.config_file, sheet_name='Vehicle_Config')
        vehicles = df[df['Is_Active'] == 'Yes'].to_dict('records')
        
        # Convert to proper format
        for vehicle in vehicles:
            if str(vehicle['Count']).lower() == 'unlimited':
                vehicle['count'] = float('inf')
            else:
                vehicle['count'] = int(vehicle['Count'])
            
            if str(vehicle['Max_Radius_KM']).lower() == 'unlimited':
                vehicle['max_radius'] = float('inf')
            else:
                vehicle['max_radius'] = float(vehicle['Max_Radius_KM'])
            
            vehicle['type'] = vehicle['Vehicle_Type']
            vehicle['capacity'] = int(vehicle['Capacity_Shipments'])
            vehicle['max_trip_time'] = int(vehicle['Max_Trip_Time_Minutes'])
        
        return vehicles

    def get_settings(self):
        df = pd.read_excel(self.config_file, sheet_name='Company_Settings')
        settings = {}
        for _, row in df.iterrows():
            settings[row['Setting_Name']] = row['Setting_Value']
        return settings

    def add_store(self, name, latitude, longitude):
        df = pd.read_excel(self.config_file, sheet_name='Store_Locations')
        new_id = df['Store_ID'].max() + 1 if not df.empty else 1
        
        new_store = pd.DataFrame([{
            'Store_ID': new_id,
            'Store_Name': name,
            'Latitude': float(latitude),
            'Longitude': float(longitude),
            'Is_Active': 'Yes'
        }])
        
        df = pd.concat([df, new_store], ignore_index=True)
        
        with pd.ExcelWriter(self.config_file, mode='a', if_sheet_exists='replace') as writer:
            df.to_excel(writer, sheet_name='Store_Locations', index=False)
        
        return new_id

    def add_vehicle(self, store_id, vehicle_type, count, capacity, max_radius, max_trip_time=480):
        try:
            df = pd.read_excel(self.config_file, sheet_name='Vehicle_Config')
            # Check if Store_ID column exists
            if 'Store_ID' not in df.columns:
                df.insert(1, 'Store_ID', 1)  # Add Store_ID column with default value 1
            new_id = df['Vehicle_ID'].max() + 1 if not df.empty else 1
        except (ValueError, FileNotFoundError):
            df = pd.DataFrame(columns=['Vehicle_ID', 'Store_ID', 'Vehicle_Type', 'Count', 'Capacity_Shipments', 'Max_Radius_KM', 'Max_Trip_Time_Minutes', 'Is_Active'])
            new_id = 1
        
        count_val = "Unlimited" if count == 'unlimited' or count == float('inf') else int(count)
        radius_val = "Unlimited" if max_radius == 'unlimited' or max_radius == float('inf') else float(max_radius)
        
        new_vehicle = pd.DataFrame([{
            'Vehicle_ID': new_id,
            'Store_ID': int(store_id),
            'Vehicle_Type': vehicle_type,
            'Count': count_val,
            'Capacity_Shipments': int(capacity),
            'Max_Radius_KM': radius_val,
            'Max_Trip_Time_Minutes': int(max_trip_time),
            'Is_Active': 'Yes'
        }])
        
        df = pd.concat([df, new_vehicle], ignore_index=True)
        
        with pd.ExcelWriter(self.config_file, mode='a', if_sheet_exists='replace') as writer:
            df.to_excel(writer, sheet_name='Vehicle_Config', index=False)
        
        return new_id

    def delete_store(self, store_id):
        df = pd.read_excel(self.config_file, sheet_name='Store_Locations')
        df.loc[df['Store_ID'] == int(store_id), 'Is_Active'] = 'No'
        
        with pd.ExcelWriter(self.config_file, mode='a', if_sheet_exists='replace') as writer:
            df.to_excel(writer, sheet_name='Store_Locations', index=False)

    def delete_vehicle(self, vehicle_id):
        df = pd.read_excel(self.config_file, sheet_name='Vehicle_Config')
        df.loc[df['Vehicle_ID'] == int(vehicle_id), 'Is_Active'] = 'No'
        
        with pd.ExcelWriter(self.config_file, mode='a', if_sheet_exists='replace') as writer:
            df.to_excel(writer, sheet_name='Vehicle_Config', index=False)

    def get_deliveries(self):
        try:
            df = pd.read_excel(self.config_file, sheet_name='Delivery_Locations')
            return df[df['Is_Active'] == 'Yes'].to_dict('records')
        except (ValueError, FileNotFoundError):
            return []

    def add_delivery(self, store_id, customer_name, address, latitude, longitude, timeslot="09:00-12:00"):
        try:
            df = pd.read_excel(self.config_file, sheet_name='Delivery_Locations')
            new_id = df['Delivery_ID'].max() + 1 if not df.empty else 1001
        except (ValueError, FileNotFoundError):
            # Create empty dataframe if sheet doesn't exist
            df = pd.DataFrame(columns=['Delivery_ID', 'Store_ID', 'Customer_Name', 'Address', 'Latitude', 'Longitude', 'Delivery_Timeslot', 'Is_Active'])
            new_id = 1001
        
        new_delivery = pd.DataFrame([{
            'Delivery_ID': new_id,
            'Store_ID': int(store_id),
            'Customer_Name': customer_name,
            'Address': address,
            'Latitude': float(latitude),
            'Longitude': float(longitude),
            'Delivery_Timeslot': timeslot,
            'Is_Active': 'Yes'
        }])
        
        df = pd.concat([df, new_delivery], ignore_index=True)
        
        with pd.ExcelWriter(self.config_file, mode='a', if_sheet_exists='replace') as writer:
            df.to_excel(writer, sheet_name='Delivery_Locations', index=False)
        
        return new_id

    def delete_delivery(self, delivery_id):
        try:
            df = pd.read_excel(self.config_file, sheet_name='Delivery_Locations')
            df.loc[df['Delivery_ID'] == int(delivery_id), 'Is_Active'] = 'No'
            
            with pd.ExcelWriter(self.config_file, mode='a', if_sheet_exists='replace') as writer:
                df.to_excel(writer, sheet_name='Delivery_Locations', index=False)
        except (ValueError, FileNotFoundError):
            pass