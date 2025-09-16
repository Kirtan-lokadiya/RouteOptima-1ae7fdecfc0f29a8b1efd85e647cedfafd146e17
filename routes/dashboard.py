from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, session
from models.config_manager import ConfigManager
import json

dashboard_bp = Blueprint('dashboard', __name__)

@dashboard_bp.route('/dashboard')
def dashboard():
    config = ConfigManager()
    stores = config.get_stores()
    vehicles = config.get_vehicles()
    deliveries = config.get_deliveries()
    settings = config.get_settings()
    
    return render_template('dashboard.html', stores=stores, vehicles=vehicles, deliveries=deliveries, settings=settings)

@dashboard_bp.route('/api/stores', methods=['GET', 'POST'])
def manage_stores():
    config = ConfigManager()
    
    if request.method == 'POST':
        data = request.get_json()
        store_id = config.add_store(
            name=data['name'],
            latitude=data['latitude'],
            longitude=data['longitude']
        )
        return jsonify({'success': True, 'store_id': store_id})
    
    stores = config.get_stores()
    return jsonify(stores)

@dashboard_bp.route('/api/stores/<store_id>', methods=['DELETE'])
def delete_store(store_id):
    config = ConfigManager()
    config.delete_store(store_id)
    return jsonify({'success': True})

@dashboard_bp.route('/api/vehicles', methods=['GET', 'POST'])
def manage_vehicles():
    config = ConfigManager()
    
    if request.method == 'POST':
        data = request.get_json()
        vehicle_id = config.add_vehicle(
            store_id=data['store_id'],
            vehicle_type=data['type'],
            count=data['count'],
            capacity=data['capacity'],
            max_radius=data['max_radius'],
            max_trip_time=data.get('max_trip_time', 480)
        )
        return jsonify({'success': True, 'vehicle_id': vehicle_id})
    
    vehicles = config.get_vehicles()
    return jsonify(vehicles)

@dashboard_bp.route('/api/vehicles/<vehicle_id>', methods=['DELETE'])
def delete_vehicle(vehicle_id):
    config = ConfigManager()
    config.delete_vehicle(vehicle_id)
    return jsonify({'success': True})

@dashboard_bp.route('/select_store')
def select_store():
    config = ConfigManager()
    stores = config.get_stores()
    return render_template('select_store.html', stores=stores)

@dashboard_bp.route('/RouteOptima_Config.xlsx')
def download_config():
    from flask import send_file
    config = ConfigManager()  # This ensures file exists
    return send_file('RouteOptima_Config.xlsx', as_attachment=True)

@dashboard_bp.route('/api/deliveries', methods=['GET', 'POST'])
def manage_deliveries():
    config = ConfigManager()
    
    if request.method == 'POST':
        data = request.get_json()
        delivery_id = config.add_delivery(
            store_id=data['store_id'],
            customer_name=data['customer_name'],
            address=data['address'],
            latitude=data['latitude'],
            longitude=data['longitude'],
            timeslot=data.get('timeslot', '09:00-12:00')
        )
        return jsonify({'success': True, 'delivery_id': delivery_id})
    
    deliveries = config.get_deliveries()
    return jsonify(deliveries)

@dashboard_bp.route('/api/deliveries/<delivery_id>', methods=['DELETE'])
def delete_delivery(delivery_id):
    config = ConfigManager()
    config.delete_delivery(delivery_id)
    return jsonify({'success': True})

@dashboard_bp.route('/generate-shipment-template/<store_id>')
def generate_shipment_template(store_id):
    from flask import send_file
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    
    config = ConfigManager()
    stores = config.get_stores()
    store = next((s for s in stores if s['Store_ID'] == int(store_id)), None)
    
    if not store:
        return "Store not found", 404
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Shipments_Data"
    
    # Headers
    headers = ['Shipment ID', 'Latitude', 'Longitude', 'Delivery Timeslot', 'Customer Name', 'Address']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    
    # Sample data around store location
    store_lat, store_lon = store['Latitude'], store['Longitude']
    sample_data = [
        [1001, store_lat + 0.01, store_lon + 0.01, "09:00-12:00", "John Doe", "123 Main St"],
        [1002, store_lat - 0.01, store_lon + 0.02, "09:00-12:00", "Jane Smith", "456 Oak Ave"],
        [1003, store_lat + 0.02, store_lon - 0.01, "12:00-15:00", "Bob Johnson", "789 Pine Rd"],
        [1004, store_lat - 0.02, store_lon - 0.02, "15:00-18:00", "Alice Brown", "321 Elm St"],
        [1005, store_lat + 0.015, store_lon + 0.015, "18:00-21:00", "Charlie Wilson", "654 Maple Dr"]
    ]
    
    for row, data_row in enumerate(sample_data, 2):
        for col, value in enumerate(data_row, 1):
            ws.cell(row=row, column=col, value=value)
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        as_attachment=True,
        download_name=f'{store["Store_Name"]}_Shipments_Template.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

