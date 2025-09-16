from flask import Blueprint, send_file, request, jsonify
import pandas as pd
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from decorators import login_required

excel_bp = Blueprint('excel', __name__)

@excel_bp.route('/api/generate-template', methods=['POST'])
@login_required
def generate_excel_template():
    data = request.get_json()
    store_name = data.get('store_name', 'Store')
    store_lat = data.get('latitude', 0.0)
    store_lon = data.get('longitude', 0.0)
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Shipments_Data"
    
    # Headers
    headers = [
        'Shipment ID', 'Latitude', 'Longitude', 'Delivery Timeslot',
        'Customer Name', 'Address', 'Phone', 'Priority'
    ]
    
    # Style headers
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # Sample data with store coordinates
    sample_data = [
        [1001, store_lat + 0.01, store_lon + 0.01, "09:00-12:00", "John Doe", "123 Main St", "+1234567890", "High"],
        [1002, store_lat - 0.01, store_lon + 0.02, "09:00-12:00", "Jane Smith", "456 Oak Ave", "+1234567891", "Medium"],
        [1003, store_lat + 0.02, store_lon - 0.01, "12:00-15:00", "Bob Johnson", "789 Pine Rd", "+1234567892", "Low"],
        [1004, store_lat - 0.02, store_lon - 0.02, "15:00-18:00", "Alice Brown", "321 Elm St", "+1234567893", "High"],
        [1005, store_lat + 0.015, store_lon + 0.015, "18:00-21:00", "Charlie Wilson", "654 Maple Dr", "+1234567894", "Medium"]
    ]
    
    for row, data_row in enumerate(sample_data, 2):
        for col, value in enumerate(data_row, 1):
            ws.cell(row=row, column=col, value=value)
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        as_attachment=True,
        download_name=f'{store_name}_shipments_template.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@excel_bp.route('/download-template')
@login_required
def download_template_page():
    return '''
    <!DOCTYPE html>
    <html>
    <head>
        <title>Download Template</title>
        <style>
            body { font-family: Arial, sans-serif; max-width: 600px; margin: 50px auto; padding: 20px; }
            .form-group { margin-bottom: 15px; }
            label { display: block; margin-bottom: 5px; font-weight: bold; }
            input { width: 100%; padding: 8px; border: 1px solid #ddd; border-radius: 4px; }
            button { background: #007bff; color: white; padding: 10px 20px; border: none; border-radius: 4px; cursor: pointer; }
            button:hover { background: #0056b3; }
        </style>
    </head>
    <body>
        <h2>Generate Excel Template</h2>
        <form id="templateForm">
            <div class="form-group">
                <label>Store Name:</label>
                <input type="text" id="storeName" required>
            </div>
            <div class="form-group">
                <label>Store Latitude:</label>
                <input type="number" step="any" id="latitude" required>
            </div>
            <div class="form-group">
                <label>Store Longitude:</label>
                <input type="number" step="any" id="longitude" required>
            </div>
            <button type="submit">Generate & Download Template</button>
        </form>
        
        <script>
            document.getElementById('templateForm').addEventListener('submit', function(e) {
                e.preventDefault();
                
                const data = {
                    store_name: document.getElementById('storeName').value,
                    latitude: parseFloat(document.getElementById('latitude').value),
                    longitude: parseFloat(document.getElementById('longitude').value)
                };
                
                fetch('/api/generate-template', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify(data)
                })
                .then(response => response.blob())
                .then(blob => {
                    const url = window.URL.createObjectURL(blob);
                    const a = document.createElement('a');
                    a.href = url;
                    a.download = data.store_name + '_shipments_template.xlsx';
                    a.click();
                    window.URL.revokeObjectURL(url);
                });
            });
        </script>
    </body>
    </html>
    '''